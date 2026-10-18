"""生成一份"故意写错"的示例文档，用来演示 docxaudit 能查出什么。

    python examples/make_sample.py
    python -m docxaudit check examples/sample-broken.docx --no-fail
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

BODY = (
    "复合材料电线杆在岛礁环境下需要同时满足抗风、耐腐蚀与运输限制，"
    "因此采用变截面薄壁结构，并在杆顶预留金具安装接口。"
)


def set_east_asia(run, font: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def body(doc, text: str = BODY, font: str = "宋体", pt: int = 22):
    p = doc.add_paragraph()
    set_east_asia(p.add_run(text), font)
    p.paragraph_format.line_spacing = Pt(pt)
    p.paragraph_format.first_line_indent = Pt(24)
    return p


def build(path: Path) -> None:
    doc = Document()

    doc.add_heading("第一章 绪论", 1)
    body(doc)
    body(doc)

    # 故意跳级：1 级之后直接 3 级
    doc.add_heading("1.1.1 研究背景", 3)

    # 正文提到图 6，但全文没有图 6
    body(doc, "如图 6 所示，杆顶位移随高度增加而迅速增大，这一结论与现场实测一致。")

    doc.add_paragraph("图 1  杆身截面、壁厚与铺层方案")
    # 故意跳号：图 1 之后直接图 3
    doc.add_paragraph("图 3  弯矩与应力沿杆长分布")

    # 故意换字体
    body(doc, "本节数据来自三次重复试验，取均值并给出标准差，用于后续校核。", font="黑体")

    # 行距故意不同
    p = doc.add_paragraph()
    set_east_asia(p.add_run("上面这一段的行距与正文其余部分不一致，打印出来一眼就能看出来。"), "宋体")
    p.paragraph_format.line_spacing = Pt(30)

    body(doc, "本文引用了相关标准与文献[1]，并在此基础上补充了成本核算部分。")

    doc.add_heading("参考文献", 1)
    doc.add_paragraph("[1] GB/T 41491—2022. 配网用复合材料杆塔[S].")
    doc.add_paragraph("[2] 王耀先. 复合材料力学与结构设计[M]. 上海: 华东理工大学出版社, 2012.")
    doc.add_paragraph("[3] 沈观林, 胡更开, 刘彬. 复合材料力学[M]. 北京: 清华大学出版社, 2013.")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    print(f"已生成：{path}")


if __name__ == "__main__":
    build(Path(__file__).with_name("sample-broken.docx"))
