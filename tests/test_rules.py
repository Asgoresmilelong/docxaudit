"""规则测试：全部用 python-docx 现场生成 .docx 夹具，不依赖任何外部样本文件。"""

from __future__ import annotations

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from docx import Document as Docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as dqn
from docx.shared import Pt

from docxaudit import audit, load_docx
from docxaudit.cli import main as cli_main

LONG = "复合材料电线杆在岛礁环境下需要同时满足抗风、耐腐蚀与运输限制，因此采用变截面设计。"


def set_east_asia(run, font: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(dqn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(dqn("w:eastAsia"), font)


class Base(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def save(self, doc, name: str = "t.docx") -> Path:
        path = self.tmp / name
        doc.save(path)
        return path

    def findings(self, path: Path, **kwargs):
        return audit(load_docx(path), **kwargs)

    @staticmethod
    def only_rule(findings, rule: str):
        return [f for f in findings if f.rule == rule]


class TestHeadingSequence(Base):
    def test_level_jump_is_error(self):
        doc = Docx()
        doc.add_heading("第一章 绪论", 1)
        doc.add_paragraph(LONG)
        doc.add_heading("1.1.1 研究背景", 3)
        got = self.only_rule(self.findings(self.save(doc)), "R1")
        self.assertTrue(any(f.severity == "error" and "跳跃" in f.message for f in got), got)

    def test_first_heading_not_level_one(self):
        doc = Docx()
        doc.add_heading("研究背景", 2)
        doc.add_paragraph(LONG)
        got = self.only_rule(self.findings(self.save(doc)), "R1")
        self.assertTrue(any(f.severity == "warn" for f in got), got)

    def test_clean_sequence(self):
        doc = Docx()
        doc.add_heading("第一章 绪论", 1)
        doc.add_heading("1.1 研究背景", 2)
        doc.add_heading("1.1.1 问题定义", 3)
        self.assertEqual(self.only_rule(self.findings(self.save(doc)), "R1"), [])


class TestCaptions(Base):
    def test_number_gap(self):
        doc = Docx()
        doc.add_paragraph("图 1 整体结构示意")
        doc.add_paragraph("图 3 截面尺寸示意")
        got = self.only_rule(self.findings(self.save(doc)), "R2")
        self.assertTrue(any("不连续" in f.message for f in got), got)

    def test_duplicate_number(self):
        doc = Docx()
        doc.add_paragraph("表 1 材料参数")
        doc.add_paragraph("表 1 荷载工况")
        got = self.only_rule(self.findings(self.save(doc)), "R2")
        self.assertTrue(any(f.severity == "error" and "重复" in f.message for f in got), got)

    def test_chapter_style_numbering(self):
        doc = Docx()
        for n in ("3-1", "3-2", "3-3"):
            doc.add_paragraph(f"图 {n} 第 {n} 张示意图")
        got = self.only_rule(self.findings(self.save(doc)), "R2")
        self.assertFalse(any("不连续" in f.message for f in got), got)


class TestCrossReference(Base):
    def test_mention_without_caption(self):
        doc = Docx()
        doc.add_paragraph("图 1 整体结构示意")
        doc.add_paragraph("如图 2 所示，杆身由三段组成，连接处采用法兰结构。")
        got = self.only_rule(self.findings(self.save(doc)), "R3")
        self.assertTrue(any("图 2" in f.message for f in got), got)

    def test_mention_matches(self):
        doc = Docx()
        doc.add_paragraph("图 1 整体结构示意")
        doc.add_paragraph("如图 1 所示，杆身由三段组成，连接处采用法兰结构。")
        got = self.only_rule(self.findings(self.save(doc)), "R3")
        self.assertEqual(got, [])


class TestReferences(Base):
    def _doc(self):
        doc = Docx()
        doc.add_paragraph("本文方法参考了[1]与[3]的相关工作，并在参数上做了调整。")
        doc.add_heading("参考文献", 1)
        doc.add_paragraph("[1] 张三. 复合材料杆塔设计[J]. 电网技术, 2020.")
        doc.add_paragraph("[2] 李四. 结构优化方法[M]. 北京: 科学出版社, 2019.")
        return doc

    def test_dangling_citation(self):
        got = self.only_rule(self.findings(self.save(self._doc())), "R4")
        self.assertTrue(any(f.severity == "error" and "[3]" in f.message for f in got), got)

    def test_uncited_entry_is_info(self):
        got = self.only_rule(self.findings(self.save(self._doc())), "R4")
        infos = [f for f in got if f.severity == "info"]
        self.assertTrue(any("[2]" in f.message for f in infos), got)

    def test_no_reference_section(self):
        doc = Docx()
        doc.add_paragraph(LONG)
        self.assertEqual(self.only_rule(self.findings(self.save(doc)), "R4"), [])


class TestFonts(Base):
    def _doc(self, fonts):
        doc = Docx()
        for font in fonts:
            p = doc.add_paragraph()
            set_east_asia(p.add_run(LONG), font)
        return doc

    def test_outlier_font(self):
        doc = self._doc(["宋体", "宋体", "宋体", "宋体", "黑体"])
        got = self.only_rule(self.findings(self.save(doc)), "R5")
        self.assertTrue(any("黑体" in f.message for f in got), got)

    def test_uniform_font(self):
        doc = self._doc(["宋体"] * 5)
        self.assertEqual(self.only_rule(self.findings(self.save(doc)), "R5"), [])

    def test_mixed_inside_paragraph(self):
        doc = Docx()
        p = doc.add_paragraph()
        set_east_asia(p.add_run("这一段前半部分使用宋体，"), "宋体")
        set_east_asia(p.add_run("后半部分突然变成黑体，看起来很不协调。"), "黑体")
        got = self.only_rule(self.findings(self.save(doc)), "R5")
        self.assertTrue(any("混用" in f.message for f in got), got)


class TestBodyScope(Base):
    def test_reference_entries_are_not_body_text(self):
        doc = Docx()
        for _ in range(5):
            p = doc.add_paragraph()
            set_east_asia(p.add_run(LONG), "宋体")
            p.paragraph_format.line_spacing = Pt(22)
            p.paragraph_format.first_line_indent = Pt(24)
        doc.add_heading("参考文献", 1)
        doc.add_paragraph("[1] 张三. 复合材料杆塔设计研究[J]. 电网技术, 2020, 44(3): 1-9.")
        doc.add_paragraph("[2] 李四. 结构优化方法及其工程应用[M]. 北京: 科学出版社, 2019.")
        path = self.save(doc)

        doc_model = load_docx(path)
        from docxaudit.rules import body_paragraphs

        bodies = list(body_paragraphs(doc_model))
        self.assertEqual(len(bodies), 5, [p.no for p in bodies])
        self.assertEqual(self.only_rule(self.findings(path), "R5"), [])


class TestSpacing(Base):
    def _doc(self, outliers: int = 0):
        doc = Docx()
        for i in range(5):
            p = doc.add_paragraph(LONG)
            p.paragraph_format.line_spacing = Pt(22)
            p.paragraph_format.first_line_indent = Pt(24)
        for _ in range(outliers):
            p = doc.add_paragraph(LONG)
            p.paragraph_format.line_spacing = Pt(30)
            p.paragraph_format.first_line_indent = Pt(24)
        return doc

    def test_outlier_spacing(self):
        got = self.only_rule(self.findings(self.save(self._doc(outliers=1))), "R6")
        self.assertTrue(got, got)

    def test_uniform_spacing(self):
        self.assertEqual(self.only_rule(self.findings(self.save(self._doc())), "R6"), [])


class TestLoader(Base):
    def test_order_and_table_flag(self):
        doc = Docx()
        doc.add_paragraph("第一段正文，用于验证段落顺序。")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格里的文字"
        doc.add_paragraph("第二段正文，仍然需要保持顺序。")
        parsed = load_docx(self.save(doc))

        self.assertEqual(len(parsed.paragraphs), 3)
        self.assertFalse(parsed.paragraphs[0].in_table)
        self.assertTrue(parsed.paragraphs[1].in_table)
        self.assertFalse(parsed.paragraphs[2].in_table)
        self.assertIn("表格里的文字", parsed.paragraphs[1].text)

    def test_rejects_non_docx(self):
        bad = self.tmp / "bad.docx"
        bad.write_bytes(b"not a zip")
        with self.assertRaises(Exception):
            load_docx(bad)


class TestFilters(Base):
    def test_ignore(self):
        doc = Docx()
        doc.add_heading("第一章 绪论", 1)
        doc.add_heading("1.1.1 背景", 3)
        path = self.save(doc)
        self.assertTrue(self.only_rule(self.findings(path), "R1"))
        self.assertEqual(self.only_rule(self.findings(path, ignore=["R1"]), "R1"), [])

    def test_only(self):
        doc = Docx()
        doc.add_heading("第一章 绪论", 1)
        doc.add_heading("1.1.1 背景", 3)
        got = self.findings(self.save(doc), only=["R2"])
        self.assertTrue(all(f.rule == "R2" for f in got), got)


class TestCli(Base):
    def test_check_returns_zero_with_no_fail(self):
        doc = Docx()
        doc.add_heading("第一章 绪论", 1)
        doc.add_heading("1.1.1 背景", 3)
        path = self.save(doc)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = cli_main(["check", str(path), "--no-fail"])
        self.assertEqual(code, 0)
        self.assertIn("R1", buf.getvalue())

    def test_check_fails_on_error(self):
        doc = Docx()
        doc.add_paragraph("图 1 结构示意")
        doc.add_paragraph("表 1 参数表")
        doc.add_paragraph("表 1 又一张参数表")
        path = self.save(doc)
        with contextlib.redirect_stdout(io.StringIO()):
            code = cli_main(["check", str(path)])
        self.assertEqual(code, 1)

    def test_json_output(self):
        doc = Docx()
        doc.add_paragraph("图 1 结构示意")
        path = self.save(doc)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            cli_main(["check", str(path), "--json", "--no-fail"])
        self.assertIn('"findings"', buf.getvalue())

    def test_missing_file(self):
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(cli_main(["check", str(self.tmp / "nope.docx")]), 2)

    def test_rules_listing(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            self.assertEqual(cli_main(["rules"]), 0)
        self.assertIn("R6", buf.getvalue())


if __name__ == "__main__":
    unittest.main()
